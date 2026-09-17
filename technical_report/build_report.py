from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

ROOT = Path(__file__).resolve().parent
SRC = ROOT / 'TECHNICAL_NOTE_SOURCE.md'
DOCX = ROOT / 'AERMOD_SURROGATE_TECHNICAL_NOTE_v0.1.docx'
PDF = ROOT / 'AERMOD_SURROGATE_TECHNICAL_NOTE_v0.1.pdf'
SITE_PDF = ROOT.parent / 'docs' / 'assets' / PDF.name

lines = SRC.read_text(encoding='utf-8').splitlines()

def clean_inline(s: str) -> str:
    return s.replace('**', '').replace('`', '')

def parse_blocks(lines):
    blocks = []
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s.startswith('|') and i + 1 < len(lines) and lines[i + 1].strip().startswith('|---'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [clean_inline(c.strip()) for c in lines[i].strip().strip('|').split('|')]
                if not all(set(c) <= {'-', ':'} for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(('table', rows))
            continue
        if s.startswith('# '): blocks.append(('h1', clean_inline(s[2:])))
        elif s.startswith('## '): blocks.append(('h2', clean_inline(s[3:])))
        elif s.startswith('### '): blocks.append(('h3', clean_inline(s[4:])))
        elif s.startswith('- '): blocks.append(('bullet', clean_inline(s[2:])))
        else: blocks.append(('p', clean_inline(s.replace('  ', ' '))))
        i += 1
    return blocks

blocks = parse_blocks(lines)
# DOCX

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.8)
sec.right_margin = Inches(0.8)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10.5)
styles['Title'].font.name = 'Aptos Display'
styles['Title'].font.size = Pt(24)
styles['Heading 1'].font.name = 'Aptos Display'
styles['Heading 1'].font.size = Pt(17)
styles['Heading 2'].font.name = 'Aptos Display'
styles['Heading 2'].font.size = Pt(13)

first_title = True
for kind, val in blocks:
    if kind == 'h1':
        if first_title:
            p = doc.add_paragraph(style='Title')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(val)
            first_title = False
        else:
            doc.add_heading(val, level=1)
    elif kind == 'h2':
        doc.add_heading(val, level=1)
    elif kind == 'h3':
        doc.add_heading(val, level=2)
    elif kind == 'bullet':
        doc.add_paragraph(val, style='List Bullet')
    elif kind == 'p':
        p = doc.add_paragraph(val)
        p.paragraph_format.space_after = Pt(6)
    elif kind == 'table':
        t = doc.add_table(rows=1, cols=len(val[0]))
        t.style = 'Table Grid'
        for j, cell in enumerate(val[0]):
            t.rows[0].cells[j].text = cell
        for row in val[1:]:
            cells = t.add_row().cells
            for j, cell in enumerate(row):
                cells[j].text = cell

doc.save(DOCX)
# PDF

base = getSampleStyleSheet()
body = ParagraphStyle('BodyHuman', parent=base['BodyText'], fontName='Helvetica', fontSize=9.5, leading=13, spaceAfter=6)
h1 = ParagraphStyle('H1Human', parent=base['Heading1'], fontName='Helvetica-Bold', fontSize=16, leading=19, spaceBefore=12, spaceAfter=7)
h2 = ParagraphStyle('H2Human', parent=base['Heading2'], fontName='Helvetica-Bold', fontSize=12, leading=15, spaceBefore=10, spaceAfter=5)
title = ParagraphStyle('TitleHuman', parent=base['Title'], fontName='Helvetica-Bold', fontSize=22, leading=26, alignment=TA_CENTER, spaceAfter=14)
small = ParagraphStyle('SmallHuman', parent=body, fontSize=8.3, leading=11)

story = []
first_title = True
for kind, val in blocks:
    if kind == 'h1':
        if first_title:
            story.append(Paragraph(escape(val), title))
            first_title = False
        else:
            story.append(Paragraph(escape(val), h1))
    elif kind == 'h2':
        story.append(Paragraph(escape(val), h1))
    elif kind == 'h3':
        story.append(Paragraph(escape(val), h2))
    elif kind == 'bullet':
        story.append(Paragraph('&bull; ' + escape(val), body))
    elif kind == 'p':
        txt = escape(val).replace('  ', ' ')
        story.append(Paragraph(txt, body))
    elif kind == 'table':
        pdata = [[Paragraph(escape(c), small) for c in row] for row in val]
        table = Table(pdata, repeatRows=1, hAlign='LEFT')
        table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#ECEFF1')),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('GRID', (0,0), (-1,-1), 0.35, colors.HexColor('#9AA0A6')),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('LEFTPADDING', (0,0), (-1,-1), 5),
            ('RIGHTPADDING', (0,0), (-1,-1), 5),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]))
        story.extend([table, Spacer(1, 7)])

def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont('Helvetica', 8)
    canvas.setFillColor(colors.HexColor('#666666'))
    canvas.drawString(18*mm, 12*mm, 'Atmospheric Dispersion Surrogate Research - public technical note')
    canvas.drawRightString(192*mm, 12*mm, f'{doc.page}')
    canvas.restoreState()

pdf = SimpleDocTemplate(str(PDF), pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=16*mm, bottomMargin=18*mm)
pdf.build(story, onFirstPage=footer, onLaterPages=footer)
SITE_PDF.write_bytes(PDF.read_bytes())
print(DOCX)
print(PDF)
print(SITE_PDF)